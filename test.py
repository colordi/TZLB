from docxtpl import DocxTemplate, InlineImage
from docx.shared import Mm
from pathlib import Path
import pandas as pd
import re

# 基础配置
TEMPLATE_PATH = Path("工作单模板.docx")
EXCEL_PATH = Path("2025 年调查数据.xlsx")
OUTPUT_DIR = Path("output")
IMAGES_DIR = Path("images")
MAX_IMAGES = 4
IMAGE_WIDTH_MM = 70


def ensure_dirs() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)


def to_date_series(series: pd.Series) -> pd.Series:
    dt = pd.to_datetime(series, errors="coerce")
    return dt.dt.date


def extract_number_from_name(path: Path) -> int:
    # 从文件名中提取形如 "-数字" 的序号，用于排序
    m = re.search(r"-(\d+)", path.stem)
    return int(m.group(1)) if m else 0


def find_image_paths(location_id) -> list[Path]:
    # 仅在 images 目录顶层查找以 location_id 开头的文件
    if not IMAGES_DIR.exists():
        return []
    prefix = str(location_id)
    candidates = [p for p in IMAGES_DIR.iterdir() if p.is_file() and p.name.startswith(prefix)]
    # 按文件名中的序号排序
    return sorted(candidates, key=extract_number_from_name)


def map_images_to_context(doc: DocxTemplate, image_paths: list[Path], region: str | None) -> dict:
    # 生成 img1..img4 映射，最多插入 MAX_IMAGES 张
    placeholders = ["img1", "img2", "img3", "img4"]
    ctx = {name: "" for name in placeholders}

    # 乡镇：第一个占位留空，从 img2 开始放图
    start_index = 1 if region == "乡镇" else 0

    inline_images = [InlineImage(doc, str(p), width=Mm(IMAGE_WIDTH_MM)) for p in image_paths[:MAX_IMAGES]]
    for i, img in enumerate(inline_images):
        slot = i + start_index
        if slot >= len(placeholders):
            break
        ctx[placeholders[slot]] = img

    return ctx


def main() -> None:
    if not TEMPLATE_PATH.exists():
        raise FileNotFoundError(f"找不到模板: {TEMPLATE_PATH}")
    if not EXCEL_PATH.exists():
        raise FileNotFoundError(f"找不到数据源: {EXCEL_PATH}")

    ensure_dirs()

    df = pd.read_excel(EXCEL_PATH)
    if "survey_date" in df.columns:
        df["survey_date"] = to_date_series(df["survey_date"])

    for _, row in df.iterrows():
        doc = DocxTemplate(TEMPLATE_PATH)
        context = row.to_dict()

        town_or_street = context.get("town_or_street")
        location_name = context.get("location_name")
        survey_date = context.get("survey_date")
        serial_number = context.get("serial_number")
        location_id = context.get("location_id")

        image_paths = find_image_paths(location_id)
        context.update(map_images_to_context(doc, image_paths, context.get("region")))

        doc.render(context)
        output_path = OUTPUT_DIR / f"2025林业有害生物防治工作单（{town_or_street}）-{location_name}-{survey_date}-{serial_number}.docx"
        doc.save(output_path)


if __name__ == "__main__":
    main()
