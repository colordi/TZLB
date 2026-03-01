from flask import Flask, render_template, request, jsonify, Response, send_file, redirect, url_for
from docxtpl import DocxTemplate, InlineImage
from docx.shared import Mm
from pathlib import Path
from datetime import datetime
from urllib.parse import quote
import re
import uuid
import base64
import csv
import io
import zipfile
import json
import math
from pest_db import (
    init_all_pest_dbs,
    insert_pest_record,
    insert_pest_records_batch,
    upsert_pest_records_batch_preserve_report_time,
    fetch_pest_records,
    delete_pest_record,
    PEST_DB_CONFIGS,
)

app = Flask(__name__)

# 基础配置
BASE_DIR = Path(__file__).parent
TEMPLATE_DIR = BASE_DIR / "templates"
IMAGES_DIR = BASE_DIR / "images"
TEMP_DIR = BASE_DIR / "temp_images"
MAX_IMAGES = 4
IMAGE_WIDTH_MM = 70
TONGZHOU_BOUNDARY_PATH = BASE_DIR / "data" / "通州区边界.geojson"
TONGZHOU_VILLAGES_PATH = BASE_DIR / "data" / "通州区村界.geojson"
YANGSHU_POINTS_PATH = BASE_DIR / "data" / "杨树点位.geojson"

# 害虫类型 → 模板路径映射
TEMPLATE_PATHS = {
    "春尺蠖": TEMPLATE_DIR / "春尺蠖工作单模板.docx",
    "国槐尺蠖": TEMPLATE_DIR / "国槐尺蠖工作单模板.docx",
    "其他害虫": TEMPLATE_DIR / "其他害虫工作单模板.docx",
}

# 确保目录存在
IMAGES_DIR.mkdir(parents=True, exist_ok=True)
TEMP_DIR.mkdir(parents=True, exist_ok=True)

# 初始化所有害虫数据库
init_all_pest_dbs()

def extract_number_from_name(path: Path) -> int:
    m = re.search(r"-(\d+)", path.stem)
    return int(m.group(1)) if m else 0

def find_image_paths(location_id) -> list[Path]:
    if not IMAGES_DIR.exists():
        return []
    prefix = str(location_id)
    candidates = [p for p in IMAGES_DIR.iterdir() if p.is_file() and p.name.startswith(prefix)]
    return sorted(candidates, key=extract_number_from_name)

def save_base64_images(base64_list: list[str], row_id: str) -> list[Path]:
    """将 Base64 图片数据保存为临时文件并返回路径列表"""
    paths = []
    for i, b64_data in enumerate(base64_list[:MAX_IMAGES]):
        try:
            # 提取 Base64 数据部分（去除 data:image/xxx;base64, 前缀）
            if "," in b64_data:
                b64_data = b64_data.split(",", 1)[1]
            
            img_bytes = base64.b64decode(b64_data)
            # 生成唯一文件名
            filename = f"{row_id}_{i}_{uuid.uuid4().hex[:8]}.jpg"
            filepath = TEMP_DIR / filename
            filepath.write_bytes(img_bytes)
            paths.append(filepath)
        except Exception as e:
            print(f"[警告] 保存图片时出错: {e}")
    return paths

def map_images_to_context(doc: DocxTemplate, image_paths: list[Path], region: str | None) -> dict:
    """将图片路径映射到模板上下文中的img1-img4占位符"""
    placeholders = ["img1", "img2", "img3", "img4"]
    ctx = {name: "" for name in placeholders}
    
    inline_images = [InlineImage(doc, str(p), width=Mm(IMAGE_WIDTH_MM)) for p in image_paths[:MAX_IMAGES]]
    for i, img in enumerate(inline_images):
        if i >= len(placeholders):
            break
        ctx[placeholders[i]] = img
    return ctx

def cleanup_temp_images(paths: list[Path]):
    """清理临时图片文件"""
    for p in paths:
        try:
            if p.exists():
                p.unlink()
        except Exception:
            pass


def resolve_year(survey_date: str | None) -> str:
    """从调查日期中提取年份，失败则回退到当前年份。"""
    if survey_date:
        match = re.match(r"(\d{4})", str(survey_date))
        if match:
            return match.group(1)
    return str(datetime.now().year)


def _is_xy_pair(node) -> bool:
    return (
        isinstance(node, list)
        and len(node) >= 2
        and isinstance(node[0], (int, float))
        and isinstance(node[1], (int, float))
    )


def _collect_xy_pairs(node, out: list[list[float]]):
    if _is_xy_pair(node):
        out.append([float(node[0]), float(node[1])])
        return
    if isinstance(node, list):
        for item in node:
            _collect_xy_pairs(item, out)


def _mercator_to_wgs84(x: float, y: float) -> list[float]:
    """将 EPSG:3857 坐标转换为 WGS84 经纬度。"""
    lon = x / 20037508.34 * 180.0
    lat = y / 20037508.34 * 180.0
    lat = 180.0 / math.pi * (2.0 * math.atan(math.exp(lat * math.pi / 180.0)) - math.pi / 2.0)
    return [lon, lat]


def _convert_mercator_coords(node):
    if _is_xy_pair(node):
        return _mercator_to_wgs84(float(node[0]), float(node[1]))
    if isinstance(node, list):
        return [_convert_mercator_coords(item) for item in node]
    return node


def normalize_geojson_to_wgs84(geojson_obj: dict) -> tuple[dict, bool]:
    """若检测到米制 Web Mercator 坐标，则转换为 WGS84 经纬度。"""
    features = geojson_obj.get("features", []) if geojson_obj.get("type") == "FeatureCollection" else []
    sample_xy: list[list[float]] = []
    for feature in features[:20]:
        geometry = feature.get("geometry") or {}
        coords = geometry.get("coordinates")
        if coords is not None:
            _collect_xy_pairs(coords, sample_xy)
        if len(sample_xy) >= 50:
            break

    # GeoJSON 正常经纬度范围通常在 [-180, 180], [-90, 90]
    # 超过该范围基本可判定为 EPSG:3857 米制坐标
    needs_convert = any(abs(x) > 180 or abs(y) > 90 for x, y in sample_xy)
    if not needs_convert:
        return geojson_obj, False

    for feature in features:
        geometry = feature.get("geometry") or {}
        if "coordinates" in geometry:
            geometry["coordinates"] = _convert_mercator_coords(geometry["coordinates"])
    return geojson_obj, True


def normalize_location_id(value: str | None) -> str:
    """标准化点位编号，便于跨数据源匹配。"""
    if value is None:
        return ""
    return str(value).strip().upper()


def get_chunchihuo_location_ids() -> set[str]:
    """获取春尺蠖数据库中存在的点位编号集合。"""
    records = fetch_pest_records("春尺蠖")
    return {
        normalize_location_id(record.get("location_id"))
        for record in records
        if normalize_location_id(record.get("location_id"))
    }

@app.route("/")
def index():
    return render_template("index.html")


@app.route("/records")
def records_page():
    """数据库可视化页面（统一入口，包含所有害虫类别）。"""
    return render_template("records.html")


@app.route("/map")
def map_page():
    """ESRI XYZ 瓦片地图页面。"""
    return render_template("map_test.html")


@app.route("/map-test")
def map_test_redirect():
    """兼容旧链接，重定向到正式地图路由。"""
    return redirect(url_for("map_page"), code=302)


# 字段 key → 中文标签映射
_FIELD_LABELS = {
    "survey_date": "调查日期", "region": "区域",
    "town_or_street": "乡镇/街道", "location_id": "点位编号",
    "location_name": "点位名称", "occurrence_position": "发生位置",
    "total_insect_count": "总虫口数", "damage_level": "受害程度",
    "report_time": "上报时间", "description": "详细情况描述",
    "pest_name": "害虫类别", "plot_type": "绿化性质",
    "host_plant": "危害寄主",
}

def _build_pest_definitions() -> list[dict]:
    """从 PEST_DB_CONFIGS 自动构建前端展示定义，避免手工维护。"""
    definitions = []
    for pest_type, config in PEST_DB_CONFIGS.items():
        # 展示字段 = 业务字段 + report_time（去掉 report_time 已在业务字段外）
        display_keys = list(config.fields) + ["report_time"]
        if pest_type == "春尺蠖" and "description" in display_keys and "report_time" in display_keys:
            desc_idx = display_keys.index("description")
            report_idx = display_keys.index("report_time")
            display_keys[desc_idx], display_keys[report_idx] = display_keys[report_idx], display_keys[desc_idx]
        fields = [{"key": k, "label": _FIELD_LABELS.get(k, k)} for k in display_keys]
        definitions.append({
            "pest_type": pest_type,
            "fields": fields,
        })
    return definitions

PEST_DEFINITIONS = _build_pest_definitions()

CHI_HUO_PEST_TYPES = {"春尺蠖", "国槐尺蠖"}
QITA_HAICHONG_PEST_TYPES = {"其他害虫"}
CHI_HUO_REMOVED_FIELDS = {"host_plant", "damaged_count", "web_count"}


def sanitize_chihuo_record(record: dict) -> dict:
    """移除尺蠖类不再使用的字段。"""
    return {key: value for key, value in record.items() if key not in CHI_HUO_REMOVED_FIELDS}


def sanitize_chihuo_records(records: list[dict]) -> list[dict]:
    """批量移除尺蠖类不再使用的字段。"""
    return [sanitize_chihuo_record(record) for record in records]


# 必填字段及其中文标签
REQUIRED_FIELDS = [
    ("town_or_street", "乡镇/街道"),
    ("location_id", "点位编号"),
    ("location_name", "点位名称"),
    ("survey_date", "调查日期"),
    ("description", "详细情况描述"),
]


def validate_required_fields(records: list[dict]) -> list[str]:
    """验证记录列表中的必填字段，返回错误消息列表（空列表表示通过）。"""
    errors = []
    for idx, record in enumerate(records):
        for field_key, field_label in REQUIRED_FIELDS:
            value = record.get(field_key, "")
            if not value or (isinstance(value, str) and not value.strip()):
                errors.append(f"第 {idx + 1} 条记录：{field_label} 不能为空")
    return errors


def fetch_all_records_grouped() -> list[dict]:
    """按虫种获取记录（字段不混用，前端按虫种分表展示）。"""
    groups = []
    for pest in PEST_DEFINITIONS:
        records = fetch_pest_records(pest["pest_type"])
        groups.append({
            "pest_type": pest["pest_type"],
            "fields": pest["fields"],
            "records": records,
            "total": len(records)
        })
    return groups


@app.route("/api/records", methods=["GET"])
def get_records():
    """获取数据库记录（按虫种分组返回）。"""
    try:
        groups = fetch_all_records_grouped()
        total = sum(group["total"] for group in groups)
        return jsonify({
            "success": True,
            "total": total,
            "pests": groups
        })
    except Exception as e:
        return jsonify({"success": False, "error": str(e)}), 500


@app.route("/api/records/export", methods=["GET"])
def export_records_csv():
    """导出数据库记录为 CSV（按虫种导出，字段不混用）。"""
    try:
        pest_type = request.args.get("pest_type", "").strip()
        matched = [p for p in PEST_DEFINITIONS if p["pest_type"] == pest_type]
        if not matched:
            if not pest_type and len(PEST_DEFINITIONS) == 1:
                matched = [PEST_DEFINITIONS[0]]
            else:
                return jsonify({"success": False, "error": "请指定有效的 pest_type 参数"}), 400
        pest = matched[0]
        records = fetch_pest_records(pest["pest_type"])
        fields = pest["fields"]
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow([field["label"] for field in fields])
        for row in records:
            writer.writerow([row.get(field["key"], "") for field in fields])

        csv_content = "\ufeff" + output.getvalue()
        filename = f"林业调查数据库_{pest['pest_type']}_{datetime.now().strftime('%Y%m%d')}.csv"
        safe_filename = f"records_export_{datetime.now().strftime('%Y%m%d')}.csv"
        encoded_filename = quote(filename)
        response = Response(csv_content, content_type="text/csv; charset=utf-8")
        response.headers["Content-Disposition"] = (
            f"attachment; filename=\"{safe_filename}\"; filename*=UTF-8''{encoded_filename}"
        )
        return response
    except Exception as e:
        return jsonify({"success": False, "error": str(e)}), 500


@app.route("/api/map/tongzhou-boundary", methods=["GET"])
def get_tongzhou_boundary():
    """读取并返回通州区边界 GeoJSON（自动转换到 WGS84）。"""
    try:
        if not TONGZHOU_BOUNDARY_PATH.exists():
            return jsonify({"success": False, "error": "未找到通州区边界文件"}), 404

        with TONGZHOU_BOUNDARY_PATH.open("r", encoding="utf-8") as f:
            geojson_obj = json.load(f)

        normalized_geojson, converted = normalize_geojson_to_wgs84(geojson_obj)
        feature_count = len(normalized_geojson.get("features", [])) if normalized_geojson.get("type") == "FeatureCollection" else 0

        return jsonify({
            "success": True,
            "name": "通州区边界",
            "feature_count": feature_count,
            "converted_from_3857": converted,
            "data": normalized_geojson,
        })
    except json.JSONDecodeError:
        return jsonify({"success": False, "error": "边界文件不是有效的 GeoJSON"}), 400
    except Exception as e:
        return jsonify({"success": False, "error": str(e)}), 500


@app.route("/api/map/tongzhou-villages", methods=["GET"])
def get_tongzhou_villages():
    """读取并返回通州区村界 GeoJSON（自动转换到 WGS84）。"""
    try:
        if not TONGZHOU_VILLAGES_PATH.exists():
            return jsonify({"success": False, "error": "未找到通州区村界文件"}), 404

        with TONGZHOU_VILLAGES_PATH.open("r", encoding="utf-8") as f:
            geojson_obj = json.load(f)

        normalized_geojson, converted = normalize_geojson_to_wgs84(geojson_obj)
        feature_count = len(normalized_geojson.get("features", [])) if normalized_geojson.get("type") == "FeatureCollection" else 0

        return jsonify({
            "success": True,
            "name": "通州区村界",
            "feature_count": feature_count,
            "converted_from_3857": converted,
            "data": normalized_geojson,
        })
    except json.JSONDecodeError:
        return jsonify({"success": False, "error": "村界文件不是有效的 GeoJSON"}), 400
    except Exception as e:
        return jsonify({"success": False, "error": str(e)}), 500


@app.route("/api/map/chunchihuo-points", methods=["GET"])
def get_chunchihuo_points():
    """返回仅与春尺蠖数据库记录关联的杨树点位。"""
    try:
        if not YANGSHU_POINTS_PATH.exists():
            return jsonify({"success": False, "error": "未找到杨树点位文件"}), 404

        with YANGSHU_POINTS_PATH.open("r", encoding="utf-8") as f:
            geojson_obj = json.load(f)

        if geojson_obj.get("type") != "FeatureCollection":
            return jsonify({"success": False, "error": "杨树点位文件不是 FeatureCollection"}), 400

        all_features = geojson_obj.get("features", [])
        db_location_ids = get_chunchihuo_location_ids()

        filtered_features = []
        for feature in all_features:
            properties = feature.get("properties") or {}
            point_id_raw = properties.get("编号")
            point_id = normalize_location_id(point_id_raw)
            if not point_id or point_id not in db_location_ids:
                continue

            # 统一补充前端展示所需字段
            properties["点位编号"] = point_id
            feature["properties"] = properties
            filtered_features.append(feature)

        filtered_geojson = {
            "type": "FeatureCollection",
            "features": filtered_features,
        }
        normalized_geojson, converted = normalize_geojson_to_wgs84(filtered_geojson)

        return jsonify({
            "success": True,
            "name": "春尺蠖关联点位",
            "source_total": len(all_features),
            "db_location_count": len(db_location_ids),
            "matched_count": len(filtered_features),
            "converted_from_3857": converted,
            "data": normalized_geojson,
        })
    except json.JSONDecodeError:
        return jsonify({"success": False, "error": "杨树点位文件不是有效的 GeoJSON"}), 400
    except Exception as e:
        return jsonify({"success": False, "error": str(e)}), 500

@app.route("/api/parse-text", methods=["POST"])
def parse_text():
    """使用AI解析自由文本为结构化数据"""
    try:
        from ai_parser import parse_text_with_ai
        
        data = request.json
        text = data.get("text", "").strip()
        pest_type = data.get("pest_type", "")  # 获取害虫类型
        
        if not text:
            return jsonify({"success": False, "error": "请输入需要解析的文本"}), 400

        if pest_type not in PEST_DB_CONFIGS:
            return jsonify({"success": False, "error": f"不支持的害虫类型: {pest_type}"}), 400
        
        records = parse_text_with_ai(text, pest_type)
        return jsonify({"success": True, "records": records})
        
    except ValueError as e:
        return jsonify({"success": False, "error": str(e)}), 400
    except RuntimeError as e:
        return jsonify({"success": False, "error": str(e)}), 500
    except Exception as e:
        return jsonify({"success": False, "error": f"解析失败: {str(e)}"}), 500

@app.route("/api/save-to-db", methods=["POST"])
def save_to_db():
    """手动保存数据到数据库"""
    try:
        data = request.json
        records = data.get("records", [])
        pest_type = data.get("pest_type", "")

        if pest_type not in PEST_DB_CONFIGS:
            return jsonify({
                "success": False,
                "error": f"不支持的害虫类型: {pest_type}"
            }), 400

        if not records:
            return jsonify({
                "success": False,
                "error": "没有数据需要保存"
            }), 400

        if pest_type in CHI_HUO_PEST_TYPES:
            records = sanitize_chihuo_records(records)

        # 必填字段验证
        validation_errors = validate_required_fields(records)
        if validation_errors:
            return jsonify({
                "success": False,
                "error": "请填写以下必填字段：\n" + "\n".join(validation_errors)
            }), 400

        # 统一批量保存
        success_count, fail_count, errors = insert_pest_records_batch(
            pest_type, records, replace_on_conflict=True
        )

        return jsonify({
            "success": True,
            "saved": success_count,
            "failed": fail_count,
            "total": len(records),
            "errors": errors
        })

    except Exception as e:
        return jsonify({"success": False, "error": str(e)}), 500


@app.route("/api/records/batch-edit", methods=["POST"])
def batch_edit_records():
    """数据库页面批量编辑保存（更新或插入，保留已有 report_time）。"""
    try:
        data = request.json or {}
        records = data.get("records", [])
        pest_type = data.get("pest_type", "")

        if pest_type not in PEST_DB_CONFIGS:
            return jsonify({
                "success": False,
                "error": f"不支持的害虫类型: {pest_type}"
            }), 400

        if not records:
            return jsonify({
                "success": False,
                "error": "没有可保存的变更"
            }), 400

        if pest_type in CHI_HUO_PEST_TYPES:
            records = sanitize_chihuo_records(records)

        validation_errors = validate_required_fields(records)
        if validation_errors:
            return jsonify({
                "success": False,
                "error": "请填写以下必填字段：\n" + "\n".join(validation_errors)
            }), 400

        success_count, fail_count, errors = upsert_pest_records_batch_preserve_report_time(
            pest_type, records
        )

        return jsonify({
            "success": True,
            "saved": success_count,
            "failed": fail_count,
            "total": len(records),
            "errors": errors
        })

    except Exception as e:
        return jsonify({"success": False, "error": str(e)}), 500


@app.route("/api/delete-record", methods=["POST"])
def delete_record():
    """从数据库删除单条记录"""
    try:
        data = request.json
        pest_type = data.get("pest_type")
        survey_date = data.get("survey_date")
        location_id = data.get("location_id")
        
        if not pest_type or not survey_date or not location_id:
            return jsonify({
                "success": False,
                "error": "缺少必填参数: pest_type, survey_date, location_id"
            }), 400
            
        success, msg = delete_pest_record(pest_type, survey_date, location_id)
        
        if success:
            return jsonify({"success": True, "message": msg})
        else:
            return jsonify({"success": False, "error": msg}), 400
            
    except Exception as e:
        return jsonify({"success": False, "error": str(e)}), 500


# 前端字段名 → Word模板变量名的映射
_DOC_FIELD_MAPPING = {
    "description": "detailed_description",
}


def _render_single_doc(
    template_path: Path,
    row: dict,
    idx: int,
    pest_type: str,
    task_type: str,
    task: str,
    all_temp_files: list[Path],
) -> tuple[str, io.BytesIO]:
    """渲染单条记录为 Word 文档，在内存中生成并返回 (文件名, BytesIO)。"""
    doc = DocxTemplate(template_path)

    # 构建模板上下文（排除 images 字段）
    context = {k: v for k, v in row.items() if k != "images"}
    context["pest_type"] = pest_type
    context["task_type"] = task_type
    context["task"] = task
    context["year"] = resolve_year(context.get("survey_date"))
    context["serial_number"] = str(idx + 1).zfill(3)

    # 字段映射：前端字段名 → Word模板变量名
    for frontend_key, template_key in _DOC_FIELD_MAPPING.items():
        if frontend_key in context:
            context[template_key] = context[frontend_key]

    # 处理图片：优先用户上传，回退到 location_id 匹配
    uploaded_images = row.get("images", [])
    if uploaded_images:
        row_id = f"row_{idx}_{uuid.uuid4().hex[:8]}"
        image_paths = save_base64_images(uploaded_images, row_id)
        all_temp_files.extend(image_paths)
    else:
        image_paths = find_image_paths(context.get("location_id"))

    context.update(map_images_to_context(doc, image_paths, context.get("region")))
    doc.render(context)

    # 构建文件名并写入内存缓冲区
    town = context.get("town_or_street") or "未知"
    loc = context.get("location_name") or "未命名"
    date = context.get("survey_date") or "无日期"
    sn = context.get("location_id") or context.get("serial_number") or str(uuid.uuid4())[:8]

    current_year = datetime.now().year
    filename = f"{current_year}林业有害生物防治工作单（{town}）-{loc}-{date}-{sn}.docx"
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return filename, buf


def _build_file_response(generated: list[tuple[str, io.BytesIO]]) -> Response:
    """根据文件数量构建下载响应（单文件直下 / 多文件 ZIP 打包），全部使用内存缓冲区。"""
    if len(generated) == 1:
        filename, buf = generated[0]
        encoded = quote(filename)
        response = send_file(
            buf,
            as_attachment=True,
            download_name=filename,
            mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
        response.headers["Content-Disposition"] = f"attachment; filename*=UTF-8''{encoded}"
        return response

    # 多文件在内存中打包 ZIP
    zip_buffer = io.BytesIO()
    with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zf:
        for filename, buf in generated:
            zf.writestr(filename, buf.read())
    zip_buffer.seek(0)

    current_year = datetime.now().year
    zip_filename = f"{current_year}林业工作单批量导出_{datetime.now().strftime('%Y%m%d_%H%M%S')}.zip"
    encoded = quote(zip_filename)
    response = send_file(
        zip_buffer,
        as_attachment=True,
        download_name=zip_filename,
        mimetype="application/zip",
    )
    response.headers["Content-Disposition"] = f"attachment; filename*=UTF-8''{encoded}"
    return response


@app.route("/api/generate", methods=["POST"])
def generate():
    all_temp_files: list[Path] = []
    try:
        data = request.json
        records = data.get("records", [])
        pest_type = data.get("pest_type", "")
        task_type = data.get("task_type", "")
        task = data.get("task", "")

        if pest_type not in PEST_DB_CONFIGS:
            return jsonify({"success": False, "error": f"不支持的害虫类型: {pest_type}"}), 400

        if pest_type in CHI_HUO_PEST_TYPES:
            records = sanitize_chihuo_records(records)

        validation_errors = validate_required_fields(records)
        if validation_errors:
            return jsonify({
                "success": False,
                "error": "请填写以下必填字段：\n" + "\n".join(validation_errors)
            }), 400

        template_path = TEMPLATE_PATHS.get(pest_type)
        if not template_path:
            return jsonify({"success": False, "error": f"未找到 {pest_type} 的工作单模板"}), 400

        # 逐条渲染文档（内存生成）并保存到数据库
        generated = []
        for idx, row in enumerate(records):
            filename, buf = _render_single_doc(
                template_path, row, idx, pest_type, task_type, task, all_temp_files
            )
            generated.append((filename, buf))
            insert_pest_record(pest_type, row, replace_on_conflict=True)

        cleanup_temp_images(all_temp_files)
        return _build_file_response(generated)

    except Exception as e:
        cleanup_temp_images(all_temp_files)
        return jsonify({"success": False, "error": str(e)}), 500

if __name__ == "__main__":
    app.run(debug=True, host='0.0.0.0', port=5001)
