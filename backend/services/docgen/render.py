from __future__ import annotations

import io
import uuid
import zipfile
from datetime import datetime
from pathlib import Path

from docx.shared import Mm
from docxtpl import DocxTemplate, InlineImage

from backend.config import get_settings
from backend.schemas import WorkOrderRecord
from backend.services.docgen.constants import (
    DOC_FIELD_MAPPING,
    IMAGE_WIDTH_MM,
    MAX_IMAGES,
)
from backend.services.docgen.images import resolve_record_image_paths
from backend.services.pest_registry import get_pest_config


def get_template_path(pest_type: str) -> Path:
    settings = get_settings()
    config = get_pest_config(pest_type)
    path = settings.templates_dir / config.template_filename
    if not path.exists():
        raise FileNotFoundError(f"未找到 {pest_type} 对应的模板文件")
    return path

def build_context(
    doc: DocxTemplate,
    record: WorkOrderRecord,
    pest_type: str,
    task_type: str,
    task_name: str,
    index: int,
    image_paths: list[Path],
) -> dict:
    """组装模板上下文。"""

    context = record.model_dump()
    serial_number = record.serial_number if record.serial_number is not None else index + 1
    context["pest_type"] = pest_type
    context["task_type"] = task_type
    context["task"] = task_name or task_type
    context["serial_number"] = str(serial_number).zfill(3)
    context["note"] = context.get("note") or ""
    context["damaged_plant_count"] = context.get("damaged_plant_count")
    context["web_nest_count"] = context.get("web_nest_count")
    if context["damaged_plant_count"] is None:
        context["damaged_plant_count"] = ""
    if context["web_nest_count"] is None:
        context["web_nest_count"] = ""

    for source_key, target_key in DOC_FIELD_MAPPING.items():
        context[target_key] = context.get(source_key) or context.get(target_key) or ""
    context["pest_species"] = context.get("pest_species") or pest_type
    context["host"] = context.get("host") or context.get("pest_hosts") or ""
    context["green_space_type"] = context.get("green_space_type") or context.get("plot_type") or ""
    context["tree_height"] = context.get("tree_height") or ""

    config = get_pest_config(pest_type)
    context["plot_type"] = context.get("plot_type") or ""
    for key, value in config.context_defaults.items():
        if context.get(key) in (None, ""):
            context[key] = value
    for key, value in config.context_overrides.items():
        context[key] = value

    placeholders = ["img1", "img2", "img3", "img4"]
    for placeholder in placeholders:
        context[placeholder] = ""

    for idx, image_path in enumerate(image_paths[:MAX_IMAGES]):
        context[placeholders[idx]] = InlineImage(doc, str(image_path), width=Mm(IMAGE_WIDTH_MM))

    return context


def build_output_filename(record: WorkOrderRecord, index: int, year: int | None = None) -> str:
    town = record.locality or "未知属地"
    location = record.location_name or "未命名点位"
    survey_date = record.survey_date or datetime.now().strftime("%Y-%m-%d")
    serial = record.location_id or str(index + 1).zfill(3)
    resolved_year = year if year is not None else datetime.now().year
    return f"{resolved_year}林业有害生物防治工作单（{town}）-{location}-{survey_date}-{serial}.docx"


def replace_suffix(filename: str, suffix: str) -> str:
    return f"{Path(filename).stem}{suffix}"


def ensure_template_context_complete(
    doc: DocxTemplate, context: dict, template_path: Path
) -> None:
    """确保模板中的占位字段都能在上下文中找到值。"""

    missing_variables = sorted(doc.get_undeclared_template_variables(context=context))
    if missing_variables:
        missing_list = "、".join(missing_variables)
        raise ValueError(f"模板缺少渲染字段：{missing_list}（模板：{template_path.name}）")


def ensure_template_markers_resolved(content: bytes, template_path: Path) -> None:
    """确保渲染后的文档中不存在未替换的模板标记。"""

    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        unresolved_files: list[str] = []
        for member in archive.namelist():
            if not member.startswith("word/") or not member.endswith(".xml"):
                continue
            xml_text = archive.read(member).decode("utf-8", errors="ignore")
            if "{{" in xml_text or "{%" in xml_text or "{#" in xml_text:
                unresolved_files.append(member)

    if unresolved_files:
        unresolved_list = "、".join(unresolved_files)
        raise ValueError(
            f"模板渲染后仍存在未替换占位：{unresolved_list}（模板：{template_path.name}）"
        )


def render_single_document(
    template_path: Path,
    record: WorkOrderRecord,
    pest_type: str,
    task_type: str,
    task_name: str,
    index: int,
    temp_images: list[Path],
    year: int | None = None,
) -> tuple[str, bytes]:
    """渲染单条 Word 工作单。"""

    doc = DocxTemplate(template_path)
    row_id = f"row_{index}_{uuid.uuid4().hex[:8]}"
    image_paths = resolve_record_image_paths(
        record=record,
        pest_type=pest_type,
        row_id=row_id,
        temp_images=temp_images,
    )

    context = build_context(doc, record, pest_type, task_type, task_name, index, image_paths)
    ensure_template_context_complete(doc, context, template_path)
    doc.render(context)

    buffer = io.BytesIO()
    doc.save(buffer)
    content = buffer.getvalue()
    ensure_template_markers_resolved(content, template_path)
    return build_output_filename(record, index, year), content
