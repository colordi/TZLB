from __future__ import annotations

import base64
import io
import json
import mimetypes
import re
import subprocess
import tempfile
import uuid
import zipfile
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from PIL import Image, UnidentifiedImageError
from docx.shared import Mm
from docxtpl import DocxTemplate, InlineImage

from backend.config import get_settings, normalize_output_format
from backend.logging_config import get_logger
from backend.schemas import WorkOrderGenerateRequest, WorkOrderRecord
from backend.services.pest_registry import (
    IMAGE_STRATEGY_WHITE_MOTH_AUTO,
    get_pest_config,
)


logger = get_logger(__name__)


MAX_IMAGES = 4
IMAGE_WIDTH_MM = 70
DOC_MEDIA_TYPE = "application/msword"
DOCX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
ZIP_MEDIA_TYPE = "application/zip"
DOC_CONVERT_FILTER = "MS Word 97"
SUPPORTED_IMAGE_FORMATS = {"JPEG", "PNG", "WEBP"}
SUPPORTED_IMAGE_FORMAT_LABEL = "JPEG、PNG、WebP"
DOC_FIELD_MAPPING = {
    "description": "detailed_description",
    "host_plant": "host",
    "pest_name": "pest_species",
}


@dataclass(slots=True)
class GeneratedArtifact:
    filename: str
    media_type: str
    content: bytes


def get_template_path(pest_type: str) -> Path:
    settings = get_settings()
    config = get_pest_config(pest_type)
    path = settings.templates_dir / config.template_filename
    if not path.exists():
        raise FileNotFoundError(f"未找到 {pest_type} 对应的模板文件")
    return path


def format_size_limit(size: int) -> str:
    if size % (1024 * 1024) == 0:
        return f"{size // (1024 * 1024)} MB"
    if size % 1024 == 0:
        return f"{size // 1024} KB"
    return f"{size} 字节"


def decode_base64_image(raw_value: str, index: int) -> bytes:
    encoded = raw_value.split(",", 1)[1] if "," in raw_value else raw_value
    try:
        return base64.b64decode(encoded, validate=True)
    except Exception as exc:  # noqa: BLE001
        raise ValueError(f"第 {index + 1} 张图片无法解析：不是有效的 Base64 图片数据") from exc


def ensure_image_size_limits(content: bytes, label: str, total_size: int | None = None) -> None:
    settings = get_settings()
    if len(content) > settings.workorder_image_max_bytes:
        raise ValueError(
            f"{label}超过单图大小限制（{format_size_limit(settings.workorder_image_max_bytes)}）"
        )
    if total_size is not None and total_size > settings.workorder_image_max_total_bytes:
        raise ValueError(
            f"图片总大小超过限制（{format_size_limit(settings.workorder_image_max_total_bytes)}）"
        )


def image_to_rgb(image: Image.Image) -> Image.Image:
    if image.mode in {"RGBA", "LA"} or "transparency" in image.info:
        rgba_image = image.convert("RGBA")
        background = Image.new("RGB", image.size, (255, 255, 255))
        background.paste(rgba_image, mask=rgba_image.getchannel("A"))
        return background
    return image.convert("RGB")


def resize_image_if_needed(image: Image.Image) -> Image.Image:
    settings = get_settings()
    max_dimension = settings.workorder_image_max_dimension
    longest_edge = max(image.size)
    if longest_edge <= max_dimension:
        return image

    scale = max_dimension / longest_edge
    target_size = (
        max(1, int(image.width * scale)),
        max(1, int(image.height * scale)),
    )
    resized = image.copy()
    resized.thumbnail(target_size, Image.Resampling.LANCZOS)
    return resized


def write_sanitized_image(content: bytes, output_path: Path, label: str) -> Path:
    try:
        with Image.open(io.BytesIO(content)) as image:
            image.load()
            image_format = (image.format or "").upper()
            if image_format not in SUPPORTED_IMAGE_FORMATS:
                raise ValueError(
                    f"{label}格式不支持：{image_format or '未知'}，仅支持 {SUPPORTED_IMAGE_FORMAT_LABEL}"
                )
            normalized_image = resize_image_if_needed(image_to_rgb(image))
    except UnidentifiedImageError as exc:
        raise ValueError(f"{label}不是有效图片") from exc

    output_path.parent.mkdir(parents=True, exist_ok=True)
    normalized_image.save(output_path, format="JPEG", quality=85, optimize=True)
    return output_path


def build_temp_image_path(row_id: str, index: int) -> Path:
    settings = get_settings()
    return settings.temp_dir / f"{row_id}_{index}_{uuid.uuid4().hex[:8]}.jpg"


def save_base64_images(base64_list: list[str], row_id: str) -> list[Path]:
    """将前端上传的 Base64 图片写入临时目录。"""

    paths: list[Path] = []
    total_size = 0
    for index, raw_value in enumerate(base64_list[:MAX_IMAGES]):
        label = f"第 {index + 1} 张图片"
        content = decode_base64_image(raw_value, index)
        total_size += len(content)
        ensure_image_size_limits(content, label, total_size)
        paths.append(write_sanitized_image(content, build_temp_image_path(row_id, index), label))
    return paths


def sanitize_existing_image_paths(paths: list[Path], row_id: str) -> list[Path]:
    """将磁盘图片校验并规范化为临时 JPEG 文件。"""

    sanitized_paths: list[Path] = []
    total_size = 0
    for index, path in enumerate(paths[:MAX_IMAGES]):
        label = f"图片文件 {path.name}"
        content = path.read_bytes()
        total_size += len(content)
        ensure_image_size_limits(content, label, total_size)
        sanitized_paths.append(
            write_sanitized_image(content, build_temp_image_path(row_id, index), label)
        )
    return sanitized_paths


def cleanup_temp_images(paths: list[Path]) -> None:
    for path in paths:
        try:
            if path.exists():
                path.unlink()
        except OSError:
            continue


def is_image_file(path: Path) -> bool:
    mime_type, _ = mimetypes.guess_type(path.name)
    return bool(mime_type and mime_type.startswith("image/"))


def natural_path_sort_key(path: Path) -> tuple[tuple[int, int | str], ...]:
    parts = re.split(r"(\d+)", path.stem.lower())
    return tuple(
        (0, int(part)) if part.isdigit() else (1, part)
        for part in parts
        if part
    )


def unique_existing_images(paths: list[Path]) -> list[Path]:
    images: list[Path] = []
    seen: set[str] = set()

    for path in paths:
        if not path.is_file() or not is_image_file(path):
            continue

        marker = str(path.resolve())
        if marker in seen:
            continue

        seen.add(marker)
        images.append(path)
        if len(images) >= MAX_IMAGES:
            break

    return images


def find_point_screenshot(screenshot_dir: Path, location_id: str) -> Path | None:
    normalized_location_id = (location_id or "").strip()
    if not normalized_location_id or not screenshot_dir.is_dir():
        return None

    matches = sorted(
        (
            path
            for path in screenshot_dir.iterdir()
            if path.is_file()
            and is_image_file(path)
            and path.stem.strip() == normalized_location_id
        ),
        key=natural_path_sort_key,
    )
    return matches[0] if matches else None


def find_dated_location_images(images_dir: Path, survey_date: str, location_id: str) -> list[Path]:
    normalized_location_id = (location_id or "").strip()
    normalized_date = (survey_date or "").strip()[:10]
    if not normalized_location_id or not normalized_date:
        return []

    dated_dir = images_dir / normalized_date
    if not dated_dir.is_dir():
        return []

    return sorted(
        (
            path
            for path in dated_dir.iterdir()
            if path.is_file()
            and is_image_file(path)
            and path.stem.startswith(normalized_location_id)
        ),
        key=natural_path_sort_key,
    )


def resolve_meiguobaie_image_paths(record: WorkOrderRecord) -> list[Path]:
    """按美国白蛾工作单规则自动装配图片，最多 4 张。"""

    settings = get_settings()
    image_paths: list[Path] = []

    point_screenshot = find_point_screenshot(
        settings.meiguobaie_point_screenshot_dir,
        record.location_id,
    )
    if point_screenshot is not None:
        image_paths.append(point_screenshot)

    image_paths.extend(
        find_dated_location_images(
            settings.images_dir,
            record.survey_date,
            record.location_id,
        )
    )

    return unique_existing_images(image_paths)


def resolve_record_image_paths(
    record: WorkOrderRecord,
    pest_type: str,
    row_id: str,
    temp_images: list[Path],
) -> list[Path]:
    config = get_pest_config(pest_type)
    if config.image_strategy == IMAGE_STRATEGY_WHITE_MOTH_AUTO:
        image_paths = sanitize_existing_image_paths(resolve_meiguobaie_image_paths(record), row_id)
        temp_images.extend(image_paths)
        return image_paths

    image_paths = save_base64_images(record.images, row_id) if record.images else []
    temp_images.extend(image_paths)
    return image_paths


def build_context(
    doc: DocxTemplate,
    record: WorkOrderRecord,
    pest_type: str,
    task_type: str,
    task_name: str,
    index: int,
    image_paths: list[Path],
) -> dict:
    """组装模板上下文。"""

    context = record.model_dump()
    serial_number = record.serial_number if record.serial_number is not None else index + 1
    context["pest_type"] = pest_type
    context["task_type"] = task_type
    context["task"] = task_name or task_type
    context["serial_number"] = str(serial_number).zfill(3)
    context["note"] = context.get("note") or ""
    context["damaged_plant_count"] = context.get("damaged_plant_count")
    context["web_nest_count"] = context.get("web_nest_count")
    if context["damaged_plant_count"] is None:
        context["damaged_plant_count"] = ""
    if context["web_nest_count"] is None:
        context["web_nest_count"] = ""

    for source_key, target_key in DOC_FIELD_MAPPING.items():
        context[target_key] = context.get(source_key) or context.get(target_key) or ""
    context["pest_species"] = context.get("pest_species") or pest_type
    context["host"] = context.get("host") or context.get("pest_hosts") or ""
    context["green_space_type"] = context.get("green_space_type") or context.get("plot_type") or ""
    context["tree_height"] = context.get("tree_height") or ""

    config = get_pest_config(pest_type)
    context["plot_type"] = context.get("plot_type") or ""
    for key, value in config.context_defaults.items():
        if context.get(key) in (None, ""):
            context[key] = value
    for key, value in config.context_overrides.items():
        context[key] = value

    placeholders = ["img1", "img2", "img3", "img4"]
    for placeholder in placeholders:
        context[placeholder] = ""

    for idx, image_path in enumerate(image_paths[:MAX_IMAGES]):
        context[placeholders[idx]] = InlineImage(doc, str(image_path), width=Mm(IMAGE_WIDTH_MM))

    return context


def build_output_filename(record: WorkOrderRecord, index: int) -> str:
    town = record.locality or "未知属地"
    location = record.location_name or "未命名点位"
    survey_date = record.survey_date or datetime.now().strftime("%Y-%m-%d")
    serial = record.location_id or str(index + 1).zfill(3)
    current_year = datetime.now().year
    return f"{current_year}林业有害生物防治工作单（{town}）-{location}-{survey_date}-{serial}.docx"


def replace_suffix(filename: str, suffix: str) -> str:
    return f"{Path(filename).stem}{suffix}"


def ensure_template_context_complete(
    doc: DocxTemplate, context: dict, template_path: Path
) -> None:
    """确保模板中的占位字段都能在上下文中找到值。"""

    missing_variables = sorted(doc.get_undeclared_template_variables(context=context))
    if missing_variables:
        missing_list = "、".join(missing_variables)
        raise ValueError(f"模板缺少渲染字段：{missing_list}（模板：{template_path.name}）")


def ensure_template_markers_resolved(content: bytes, template_path: Path) -> None:
    """确保渲染后的文档中不存在未替换的模板标记。"""

    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        unresolved_files: list[str] = []
        for member in archive.namelist():
            if not member.startswith("word/") or not member.endswith(".xml"):
                continue
            xml_text = archive.read(member).decode("utf-8", errors="ignore")
            if "{{" in xml_text or "{%" in xml_text or "{#" in xml_text:
                unresolved_files.append(member)

    if unresolved_files:
        unresolved_list = "、".join(unresolved_files)
        raise ValueError(
            f"模板渲染后仍存在未替换占位：{unresolved_list}（模板：{template_path.name}）"
        )


def render_single_document(
    template_path: Path,
    record: WorkOrderRecord,
    pest_type: str,
    task_type: str,
    task_name: str,
    index: int,
    temp_images: list[Path],
) -> tuple[str, bytes]:
    """渲染单条 Word 工作单。"""

    doc = DocxTemplate(template_path)
    row_id = f"row_{index}_{uuid.uuid4().hex[:8]}"
    image_paths = resolve_record_image_paths(
        record=record,
        pest_type=pest_type,
        row_id=row_id,
        temp_images=temp_images,
    )

    context = build_context(doc, record, pest_type, task_type, task_name, index, image_paths)
    ensure_template_context_complete(doc, context, template_path)
    doc.render(context)

    buffer = io.BytesIO()
    doc.save(buffer)
    content = buffer.getvalue()
    ensure_template_markers_resolved(content, template_path)
    return build_output_filename(record, index), content


def convert_docx_bytes_to_doc(filename: str, content: bytes) -> tuple[str, bytes]:
    """使用 LibreOffice 将 docx 字节流转换为 doc。"""

    settings = get_settings()
    source_filename = replace_suffix(filename, ".docx")
    target_filename = replace_suffix(filename, ".doc")

    with tempfile.TemporaryDirectory(dir=settings.temp_dir, prefix="workorder_export_") as workdir:
        workdir_path = Path(workdir)
        source_path = workdir_path / source_filename
        target_path = workdir_path / target_filename
        source_path.write_bytes(content)

        try:
            subprocess.run(
                [
                    settings.libreoffice_bin,
                    "--headless",
                    "--convert-to",
                    f"doc:{DOC_CONVERT_FILTER}",
                    "--outdir",
                    str(workdir_path),
                    str(source_path),
                ],
                check=True,
                capture_output=True,
                timeout=settings.libreoffice_timeout_seconds,
            )
        except FileNotFoundError as exc:
            raise RuntimeError(
                f"未找到 LibreOffice 命令行工具：{settings.libreoffice_bin}"
            ) from exc
        except subprocess.TimeoutExpired as exc:
            raise RuntimeError("LibreOffice 转换超时，请稍后重试。") from exc
        except subprocess.CalledProcessError as exc:
            stderr = exc.stderr.decode("utf-8", errors="ignore").strip() if exc.stderr else ""
            stdout = exc.stdout.decode("utf-8", errors="ignore").strip() if exc.stdout else ""
            detail = stderr or stdout or "未知错误"
            raise RuntimeError(f"LibreOffice 转换失败：{detail}") from exc

        if not target_path.exists():
            raise RuntimeError("LibreOffice 转换失败：未生成 .doc 文件。")

        return target_filename, target_path.read_bytes()


def build_download_artifact(filename: str, content: bytes, media_type: str) -> GeneratedArtifact:
    """返回单个下载产物。"""

    return GeneratedArtifact(
        filename=filename,
        media_type=media_type,
        content=content,
    )


def resolve_output_format(payload: WorkOrderGenerateRequest) -> str:
    settings = get_settings()
    return normalize_output_format(payload.output_format or settings.workorder_default_output_format)


def generate_workorder_artifact(payload: WorkOrderGenerateRequest) -> GeneratedArtifact:
    """生成工作单下载产物。"""

    if len(payload.records) != 1:
        raise ValueError("批量压缩导出已取消，请改为逐条导出工作单。")

    template_path = get_template_path(payload.pest_type)
    output_format = resolve_output_format(payload)
    temp_images: list[Path] = []
    record = payload.records[0]

    logger.info(
        "开始生成工作单: pest_type=%s task_type=%s output_format=%s location=%s date=%s images=%d",
        payload.pest_type,
        payload.task_type,
        output_format,
        record.location_name,
        record.survey_date,
        len(record.images),
    )

    try:
        filename, content = render_single_document(
            template_path=template_path,
            record=record,
            pest_type=payload.pest_type,
            task_type=payload.task_type,
            task_name=payload.task,
            index=0,
            temp_images=temp_images,
        )
        if output_format == "doc":
            filename, content = convert_docx_bytes_to_doc(filename, content)
            logger.info("工作单生成完成: %s (doc)", filename)
            return build_download_artifact(filename, content, DOC_MEDIA_TYPE)
        if output_format == "docx":
            logger.info("工作单生成完成: %s (docx)", filename)
            return build_download_artifact(filename, content, DOCX_MEDIA_TYPE)
        raise ValueError("输出格式只能是 doc 或 docx")
    finally:
        cleanup_temp_images(temp_images)


@dataclass(slots=True)
class BatchResult:
    filename: str
    content: bytes


@dataclass(slots=True)
class BatchFailure:
    index: int
    location_name: str
    location_id: str
    reason: str


def build_batch_zip_filename(success_count: int) -> str:
    current_year = datetime.now().year
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    return f"{current_year}年工作单批量导出_{timestamp}_{success_count}份.zip"


def generate_workorder_batch_artifact(payload: WorkOrderGenerateRequest) -> GeneratedArtifact:
    """批量生成工作单并打包为 zip。

    单条失败不影响整体流程，成功文件放入 zip，失败记录写入 失败记录.json。
    如果全部失败，则抛出 ValueError 统一返回 400。
    """

    template_path = get_template_path(payload.pest_type)
    output_format = resolve_output_format(payload)
    successes: list[BatchResult] = []
    failures: list[BatchFailure] = []

    logger.info(
        "开始批量生成工作单: pest_type=%s task_type=%s output_format=%s count=%d",
        payload.pest_type,
        payload.task_type,
        output_format,
        len(payload.records),
    )

    for index, record in enumerate(payload.records):
        temp_images: list[Path] = []
        try:
            filename, content = render_single_document(
                template_path=template_path,
                record=record,
                pest_type=payload.pest_type,
                task_type=payload.task_type,
                task_name=payload.task,
                index=index,
                temp_images=temp_images,
            )
            if output_format == "doc":
                filename, content = convert_docx_bytes_to_doc(filename, content)
            successes.append(BatchResult(filename=filename, content=content))
        except Exception as exc:  # noqa: BLE001
            reason = str(exc)
            logger.warning(
                "第 %d 条工作单生成失败: location=%s reason=%s",
                index + 1,
                record.location_name,
                reason,
            )
            failures.append(
                BatchFailure(
                    index=index + 1,
                    location_name=record.location_name or "未命名点位",
                    location_id=record.location_id or "",
                    reason=reason,
                )
            )
        finally:
            cleanup_temp_images(temp_images)

    if not successes:
        failure_details = "；".join(
            f"第 {failure.index} 条（{failure.location_name}）: {failure.reason}"
            for failure in failures
        )
        raise ValueError(f"批量导出全部失败：{failure_details}")

    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for result in successes:
            archive.writestr(f"工作单/{result.filename}", result.content)

        if failures:
            failures_json = io.StringIO()
            json.dump(
                {
                    "total": len(payload.records),
                    "success_count": len(successes),
                    "failure_count": len(failures),
                    "failures": [
                        {
                            "index": failure.index,
                            "location_name": failure.location_name,
                            "location_id": failure.location_id,
                            "reason": failure.reason,
                        }
                        for failure in failures
                    ],
                },
                failures_json,
                ensure_ascii=False,
                indent=2,
            )
            archive.writestr("失败记录.json", failures_json.getvalue())

    logger.info(
        "批量生成工作单完成: success=%d failure=%d",
        len(successes),
        len(failures),
    )

    return GeneratedArtifact(
        filename=build_batch_zip_filename(len(successes)),
        media_type=ZIP_MEDIA_TYPE,
        content=buffer.getvalue(),
    )
